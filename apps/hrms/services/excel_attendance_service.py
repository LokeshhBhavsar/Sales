from datetime import datetime

from django.db import transaction

from openpyxl import load_workbook

from apps.hrms.models.employee import Employee
from apps.hrms.models.attendance_import_row import (
    AttendanceImportRow
)

from apps.hrms.services.attendance_service import (
    AttendanceService
)


class ExcelAttendanceService:

    @staticmethod
    @transaction.atomic
    def process_import(attendance_import):

        attendance_import.status = "processing"
        attendance_import.save(
            update_fields=["status"]
        )

        try:

            workbook = load_workbook(
                attendance_import.file.path,
                data_only=True
            )

            worksheet = workbook.active

            rows = list(
                worksheet.iter_rows(
                    values_only=True
                )
            )

            if not rows:

                attendance_import.status = "failed"
                attendance_import.error_message = (
                    "Excel file is empty."
                )
                attendance_import.save()

                return attendance_import

            headers = [
                str(value).strip().lower()
                if value is not None
                else ""
                for value in rows[0]
            ]

            required_headers = [
                "employee_code",
                "date",
                "check_in",
                "check_out",
            ]

            missing_headers = [
                header
                for header in required_headers
                if header not in headers
            ]

            if missing_headers:

                attendance_import.status = "failed"
                attendance_import.error_message = (
                    "Missing columns: "
                    + ", ".join(missing_headers)
                )
                attendance_import.save()

                return attendance_import

            header_map = {
                header: index
                for index, header in enumerate(headers)
            }

            total = 0
            success = 0
            failed = 0

            for row_number, row in enumerate(
                rows[1:],
                start=2
            ):

                total += 1

                employee_code = row[
                    header_map["employee_code"]
                ]

                attendance_date = row[
                    header_map["date"]
                ]

                check_in = row[
                    header_map["check_in"]
                ]

                check_out = row[
                    header_map["check_out"]
                ]

                import_row = AttendanceImportRow.objects.create(
                    attendance_import=attendance_import,
                    row_number=row_number,
                    employee_code=str(
                        employee_code
                    ).strip()
                    if employee_code is not None
                    else "",
                )

                try:

                    if not employee_code:

                        raise ValueError(
                            "Employee code is required."
                        )

                    employee = Employee.objects.get(
                        employee_code=str(
                            employee_code
                        ).strip(),
                        company=attendance_import.company,
                        plant=attendance_import.plant,
                    )

                    if not attendance_date:

                        raise ValueError(
                            "Attendance date is required."
                        )

                    if isinstance(
                        attendance_date,
                        datetime
                    ):

                        attendance_date = (
                            attendance_date.date()
                        )

                    if isinstance(
                        check_in,
                        datetime
                    ):
                        check_in_datetime = check_in

                    else:
                        check_in_datetime = None

                    if isinstance(
                        check_out,
                        datetime
                    ):
                        check_out_datetime = check_out

                    else:
                        check_out_datetime = None

                    import_row.employee = employee
                    import_row.attendance_date = (
                        attendance_date
                    )

                    import_row.check_in = (
                        check_in_datetime
                    )

                    import_row.check_out = (
                        check_out_datetime
                    )

                    import_row.save()

                    AttendanceService.create_or_update_attendance(
                        employee=employee,
                        attendance_date=attendance_date,
                        check_in=check_in_datetime,
                        check_out=check_out_datetime,
                        source="excel"
                    )

                    import_row.status = "success"
                    import_row.save(
                        update_fields=["status"]
                    )

                    success += 1

                except Exception as exc:

                    import_row.status = "failed"
                    import_row.error_message = str(exc)

                    import_row.save(
                        update_fields=[
                            "status",
                            "error_message"
                        ]
                    )

                    failed += 1

            attendance_import.total_rows = total
            attendance_import.successful_rows = success
            attendance_import.failed_rows = failed

            if failed == 0:

                attendance_import.status = "completed"

            elif success > 0:

                attendance_import.status = (
                    "partially_completed"
                )

            else:

                attendance_import.status = "failed"

            attendance_import.save()

            return attendance_import

        except Exception as exc:

            attendance_import.status = "failed"
            attendance_import.error_message = str(exc)

            attendance_import.save()

            return attendance_import